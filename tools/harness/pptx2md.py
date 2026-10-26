"""Extract structured content from PPTX files: text, tables, images, speaker notes."""
from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu


@dataclass
class SlideBlock:
    index: int
    title: str = ""
    text_lines: list[str] = field(default_factory=list)
    table_rows: list[list[str]] = field(default_factory=list)
    images: list[bytes] = field(default_factory=list)
    image_exts: list[str] = field(default_factory=list)
    notes: str = ""


def extract_pptx(path: Path) -> tuple[list[SlideBlock], dict]:
    prs = Presentation(str(path))
    slides: list[SlideBlock] = []
    meta = {"slide_count": 0, "figures": 0}

    for idx, slide in enumerate(prs.slides, 1):
        block = SlideBlock(index=idx)

        if slide.shapes.title is not None:
            block.title = slide.shapes.title.text.strip()

        for shape in slide.shapes:
            if shape.shape_type == 13:  # picture
                try:
                    img_bytes = shape.image.blob
                    ext = shape.image.ext or "png"
                    block.images.append(img_bytes)
                    block.image_exts.append(ext)
                    meta["figures"] += 1
                except Exception:
                    pass

            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        block.text_lines.append(text)

            if shape.has_table:
                table = shape.table
                rows = []
                for row in table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                block.table_rows = rows

        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                block.notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass

        slides.append(block)

    meta["slide_count"] = len(slides)
    return slides, meta
